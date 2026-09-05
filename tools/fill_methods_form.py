#!/usr/bin/env python3
"""Fill the challenge's methods-description form from a JSON answers file.

The organisers' template (``methods_description_form.xlsx``) has one sheet per track, the
questions in column A and an empty answer column B. This tool copies the template and
writes the answers — nothing else changes, so the panel receives the file it published,
filled in. Case-agnostic: the answers live beside the submission they describe.

Answers file: a JSON object mapping the **row number** of the question (as the template
lays it out — ``7`` is *Team name*, ``18`` the abstract on the Track 1 sheet) to the answer
text. Keys are strings because JSON demands it.

Usage::

    uv run python tools/fill_methods_form.py --sheet "Track 1 methods" \\
        --answers submissions/track1/iteration-1/methods-form.json \\
        --out submissions/track1/iteration-1/team_track1_methods_form.xlsx
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment

TEMPLATE = Path("ressources/hackathon-rules/methods_description_form.xlsx")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--template", type=Path, default=TEMPLATE)
    parser.add_argument("--sheet", required=True, help='e.g. "Track 1 methods"')
    parser.add_argument("--answers", type=Path, required=True, help="JSON: row number → answer")
    parser.add_argument("--out", type=Path, required=True)
    return parser.parse_args()


def fill(template: Path, sheet: str, answers: dict[str, str], out: Path) -> list[tuple[int, str]]:
    """Write each answer into column B of its row; return (row, question) for what was filled."""
    book = openpyxl.load_workbook(template)
    page = book[sheet]
    written = []
    for key, text in answers.items():
        row = int(key)
        question = page.cell(row=row, column=1).value
        if not question:
            raise ValueError(f"row {row} has no question in column A of sheet {sheet!r}")
        cell = page.cell(row=row, column=2)
        cell.value = text
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        written.append((row, str(question)))
    out.parent.mkdir(parents=True, exist_ok=True)
    book.save(out)
    return written


def main() -> None:
    args = parse_args()
    answers = json.loads(args.answers.read_text(encoding="utf-8"))
    written = fill(args.template, args.sheet, answers, args.out)
    for row, question in written:
        print(f"row {row:>2}  {question[:70]}")
    print(f"wrote {args.out}  ({len(written)} answers on sheet {args.sheet!r})")


if __name__ == "__main__":
    main()
