"""Reproduce audit artifacts: uv run python ressources/track2/it5/evidence_checks/check_cached_evidence.py.

One-off audit, not a maintained production entry point. Source layouts and
candidate names are explicit; target annotations are not BUB1B rescue evidence.
"""

import csv
import hashlib
import json
import math
from collections import Counter
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent
RAW = Path("build/it5/r2/raw")
inputs = []


def record(path):
    inputs.append({"path": str(path), "sha256": hashlib.sha256(path.read_bytes()).hexdigest()})


def table(name, rows):
    with (OUT / name).open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), delimiter="\t")
        writer.writeheader()
        writer.writerows(rows)


labels_path = Path("ressources/track2/bub1b_published_alleles.tsv")
record(labels_path)
labels = list(csv.DictReader(labels_path.open(), delimiter="\t"))
primary = Counter(r["class"] for r in labels if r["primary"] == "yes")
assert primary == {"abundance-lowering": 5, "stable": 3}

path = RAW / "mitchell/MOESM3.csv"
record(path)
with path.open(encoding="cp1252", newline="") as handle:
    reader = csv.reader(handle)
    header = next(reader)
    bub = [r for r in reader if r[0] == "BUB1B"]
mitchell_compounds = len(header) - 2
assert len(bub) == 1
conditions = []
for drug in ("Gliclazide", "Migalastat"):
    indices = [i for i, h in enumerate(header) if drug.lower() in h.lower()]
    assert len(indices) == 1
    value = float(bub[0][indices[0]])
    conditions.append({
        "dataset": "Mitchell 2023", "compound": drug, "cell": "HCT116",
        "dose_uM": 10, "time_h": 24, "published_summary_columns": len(indices),
        "independent_replicates_recoverable": "not resolved in summary",
        "raw_BUB1B_log2FC": value, "raw_BUB1B_fold_change": 2 ** value,
        "additional_dose_time_conditions": 0,
        "citation": "https://doi.org/10.1038/s41587-022-01539-0",
    })
table("candidate_conditions.tsv", conditions)

canonical_path = Path("ressources/track2/it5/proteome_hits.tsv")
record(canonical_path)
canonical = list(csv.DictReader(canonical_path.open(), delimiter="\t"))
units = []
for condition in conditions:
    row = next(r for r in canonical if r["dataset"] == "mitchell2023"
               and r["compound"] == condition["compound"])
    assert float(row["raw_log2fc"]) == condition["raw_BUB1B_log2FC"]
    units.append({
        "compound": row["compound"], "raw_log2fc": row["raw_log2fc"],
        "raw_fold_change": condition["raw_BUB1B_fold_change"],
        "proliferation_score": row["proliferation_score"],
        "residual": row["residual"], "robust_z": row["robust_z"],
        "rank": row["rank"], "sensitivity_z": row["sensitivity_z"],
        "sensitivity_rank": row["sensitivity_rank"],
        "unsupported_narrative_value": "1.47" if row["compound"] == "Gliclazide" else "1.44",
        "unsupported_value_is_demonstrated_EB_z": False,
    })
table("canonical_units_audit.tsv", units)

metadata_matches = []
for path in sorted((RAW / "mitchell").glob("*.xlsx")):
    record(path)
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet in book:
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
            matches = [d for d in ("gliclazide", "migalastat")
                       if any(d in str(v).lower() for v in row)]
            for drug in matches:
                metadata_matches.append({
                    "file": str(path), "sheet": sheet.title, "row_1based": row_number,
                    "compound": drug, "row_values_json": json.dumps(row, default=str),
                })
    book.close()
table("candidate_workbook_matches.tsv", metadata_matches)
assert len(metadata_matches) == 2

independent = []
for name, cell in (("MOESM7.xlsx", "A549"), ("MOESM8.xlsx", "MCF-7"),
                   ("MOESM9.xlsx", "RKO")):
    path = RAW / "ptm" / name
    record(path)
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = book.worksheets[0].iter_rows(values_only=True)
    next(rows)
    next(rows)
    header = next(rows)
    candidates = [h for h in header[5:] if any(
        d in str(h).lower() for d in ("gliclazide", "migalastat"))]
    assert not candidates
    target_rows = [row for row in rows if str(row[2]) == "BUB1B"]
    assert len(target_rows) == 1
    for label, value in zip(header[5:], target_rows[0][5:], strict=True):
        independent.append({
            "dataset": "Saei 2019", "cell": cell, "condition": label,
            "BUB1B_fold_change_vs_control": value,
            "BUB1B_log2FC": math.log2(value),
            "candidate_overlap": False,
            "citation": "https://doi.org/10.1038/s41467-019-13582-8",
        })
    book.close()
assert len(independent) == 81
table("independent_BUB1B_conditions.tsv", independent)

path = OUT / "decryptE_compounds.tsv"
record(path)
compounds = list(csv.DictReader(path.open(), delimiter="\t"))
assert len(compounds) == 144
assert not any(d in json.dumps(compounds).lower()
               for d in ("gliclazide", "migalastat"))

summary = {
    "scope": "Bounded cached primary evidence, not a systematic literature review",
    "assumptions": [
        "Run from repository root with existing public caches and project uv environment",
        "Existing primary labels and their assay-context classifications are preserved",
        "Mitchell CSV first two columns are gene and accession; remaining columns are compounds",
        "Candidate workbook search uses case-insensitive literal gliclazide/migalastat names",
        "decryptE full compound and listed-alias inventory is searched, not protein measurements",
        "One summary column does not establish biological replicate count",
        "Drug target annotations, inhibition and association are not demonstrated BUB1B abundance rescue",
        "No comparable-domain panel was validated; this is not an exhaustive negative search",
    ],
    "primary_labels": dict(primary), "primary_n": 8,
    "extended_stable_n": sum(r["class"] == "stable" and r["primary"] != "yes" for r in labels),
    "ambiguous_excluded_n": sum(r["class"] == "ambiguous" for r in labels),
    "new_validated_independent_abundance_labels": 0,
    "new_validated_comparable_domain_labels": 0,
    "raw_ddg_AUROC_primary": 0.60, "reported_bootstrap95": [0.13333333333333333, 1.0],
    "calibrated_measured_prior_available": False,
    "mitchell_compound_columns": mitchell_compounds,
    "candidate_additional_dose_time_conditions": 0,
    "independent_Saei_BUB1B_columns": len(independent),
    "independent_Saei_drugs": 9, "independent_Saei_cell_lines": 3,
    "independent_Saei_replicates_per_drug_cell": 3,
    "Saei_candidate_overlap": 0, "decryptE_drugs": len(compounds),
    "decryptE_candidate_overlap": 0,
    "conclusion": "No validated expansion; candidate protein signals and stability inference remain exploratory",
    "inputs": inputs,
}
(OUT / "summary.json").write_text(json.dumps(summary, indent=2) + "\n")
print(json.dumps({k: v for k, v in summary.items() if k != "inputs"}, indent=2))
